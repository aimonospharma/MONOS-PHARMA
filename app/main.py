"""MP team — Видео контент түгээх, хэмжих платформ (monolithic FastAPI app)."""
import csv
import io
import json
from contextlib import asynccontextmanager
from datetime import datetime

from fastapi import FastAPI, Request, Form, File, UploadFile, HTTPException, Query
from fastapi.responses import (HTMLResponse, RedirectResponse, JSONResponse,
                               StreamingResponse, PlainTextResponse)
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from starlette.middleware.sessions import SessionMiddleware

from . import auth, repo, storage
from .config import (APP_NAME, APP_TAGLINE, BASE_DIR, SECRET_KEY, SESSION_COOKIE,
                     SESSION_HTTPS_ONLY,
                     SECTIONS, SECTION_SHORT, SECTION_BONUS, SECTION_OTHER,
                     ROLE_ADMIN, ROLE_VIEWER, ROLE_LABELS, tier_for, COMPANY_DOMAIN,
                     MAX_UPLOAD_MB)
from .database import init_db

@asynccontextmanager
async def lifespan(_: FastAPI):
    init_db(seed=True)      # схем үүсгэж, хоосон бол demo өгөгдөл нэмнэ
    yield


app = FastAPI(title=APP_NAME, docs_url=None, redoc_url=None, lifespan=lifespan)
app.add_middleware(SessionMiddleware, secret_key=SECRET_KEY, session_cookie=SESSION_COOKIE,
                   max_age=60 * 60 * 24 * 14, same_site="lax", https_only=SESSION_HTTPS_ONLY)
app.mount("/static", StaticFiles(directory=BASE_DIR / "static"), name="static")

templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))


# ---------------------------------------------------------------- filters
def fmt_date(value, with_time: bool = False):
    if not value:
        return "—"
    s = str(value)
    try:
        dt = datetime.fromisoformat(s.replace("T", " "))
    except ValueError:
        return s
    return dt.strftime("%Y.%m.%d %H:%M") if with_time else dt.strftime("%Y.%m.%d")


def fmt_ago(value):
    if not value:
        return "—"
    try:
        dt = datetime.fromisoformat(str(value).replace("T", " "))
    except ValueError:
        return str(value)
    delta = datetime.now() - dt
    mins = int(delta.total_seconds() // 60)
    if mins < 1:
        return "дөнгөж сая"
    if mins < 60:
        return f"{mins} мин өмнө"
    if mins < 60 * 24:
        return f"{mins // 60} цагийн өмнө"
    if delta.days < 30:
        return f"{delta.days} хоногийн өмнө"
    return dt.strftime("%Y.%m.%d")


def fmt_duration(sec):
    sec = int(sec or 0)
    if sec <= 0:
        return "—"
    h, rem = divmod(sec, 3600)
    m, s = divmod(rem, 60)
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m}:{s:02d}"


templates.env.filters["dt"] = lambda v: fmt_date(v, True)
templates.env.filters["d"] = fmt_date
templates.env.filters["ago"] = fmt_ago
templates.env.filters["dur"] = fmt_duration
templates.env.filters["size"] = storage.human_size
templates.env.filters["num"] = lambda n: f"{int(n or 0):,}".replace(",", " ")
templates.env.globals.update(
    APP_NAME=APP_NAME, APP_TAGLINE=APP_TAGLINE, SECTIONS=SECTIONS,
    SECTION_SHORT=SECTION_SHORT, SECTION_BONUS=SECTION_BONUS, SECTION_OTHER=SECTION_OTHER,
    ROLE_ADMIN=ROLE_ADMIN, ROLE_VIEWER=ROLE_VIEWER, ROLE_LABELS=ROLE_LABELS,
    COMPANY_DOMAIN=COMPANY_DOMAIN, MAX_UPLOAD_MB=MAX_UPLOAD_MB, tier_for=tier_for,
    now=datetime.now,
)

KIND_LABELS = {"video": "Видео", "image": "Зураг", "poster": "Постер", "pdf": "PDF"}
templates.env.globals["KIND_LABELS"] = KIND_LABELS


def page(request: Request, name: str, **ctx):
    user = auth.current_user(request)
    flash = request.session.pop("flash", None)
    return templates.TemplateResponse(
        request, name, {"user": user, "flash": flash, **ctx}
    )


def flash(request: Request, message: str, level: str = "ok"):
    request.session["flash"] = {"message": message, "level": level}


@app.exception_handler(HTTPException)
async def http_error(request: Request, exc: HTTPException):
    if exc.status_code == 401:
        return RedirectResponse(f"/login?next={request.url.path}", status_code=303)
    if request.url.path.startswith("/api/"):
        return JSONResponse({"error": exc.detail}, status_code=exc.status_code)
    return templates.TemplateResponse(
        request, "error.html",
        {"user": auth.current_user(request), "flash": None,
         "code": exc.status_code, "detail": exc.detail},
        status_code=exc.status_code,
    )


# =========================================================== AUTH ROUTES
@app.get("/login", response_class=HTMLResponse)
def login_form(request: Request, next: str = "/"):
    if auth.current_user(request):
        return RedirectResponse(next or "/", status_code=303)
    return page(request, "login.html", next=next, ms_accounts=auth.DEMO_MS_ACCOUNTS)


@app.post("/login")
def login_submit(request: Request, email: str = Form(""), name: str = Form(""),
                 role: str = Form(ROLE_VIEWER), next: str = Form("/")):
    try:
        user = auth.sign_in_mock(request, email, name, role)
    except ValueError as e:
        flash(request, str(e), "error")
        return RedirectResponse(f"/login?next={next}", status_code=303)
    flash(request, f"Тавтай морил, {user['name']}!")
    return RedirectResponse(next or "/", status_code=303)


@app.get("/auth/microsoft")
def microsoft_sso(request: Request, email: str = "bolormaa.b@monos.mn", next: str = "/"):
    """Demo горим: Microsoft SSO-г дуурайлган нэвтрүүлнэ.

    Production-д энд Azure AD authorize URL руу redirect хийж, /auth/microsoft/callback
    дээр код солилцоод Graph /me-ээс профайл татна (auth.py дахь MICROSOFT_SSO тэмдэглэл)."""
    user = auth.resolve_microsoft_user(request, email)
    flash(request, f"Microsoft account-аар нэвтэрлээ: {user['email']}")
    return RedirectResponse(next or "/", status_code=303)


@app.get("/logout")
def logout(request: Request):
    auth.logout(request)
    return RedirectResponse("/login", status_code=303)


# =========================================================== VIEWER SIDE
@app.get("/", response_class=HTMLResponse)
def home(request: Request):
    user = auth.require_user(request)
    if user["role"] == ROLE_ADMIN:
        return RedirectResponse("/dashboard", status_code=303)
    return RedirectResponse("/library", status_code=303)


@app.get("/library", response_class=HTMLResponse)
def library(request: Request, section: str = "", product_id: int = 0, kind: str = "",
            q: str = "", sort: str = "new"):
    user = auth.require_user(request)
    section = section if section in SECTIONS else ""
    items = repo.list_contents(section=section or None, product_id=product_id or None,
                               kind=kind or None, search=q or None, status="published")
    if sort == "popular":
        items = sorted(items, key=lambda r: r["views_total"], reverse=True)
    bonus = [i for i in items if i["section"] == SECTION_BONUS]
    other = [i for i in items if i["section"] == SECTION_OTHER]
    featured = bonus[0] if bonus else (other[0] if other else None)
    seen = {r["content_id"] for r in repo.query(
        "SELECT DISTINCT content_id FROM views WHERE user_id = ?", (user["id"],))}
    return page(request, "library.html", bonus=bonus, other=other, featured=featured,
                products=repo.list_products(), filters={"section": section, "product_id": product_id,
                                                        "kind": kind, "q": q, "sort": sort},
                seen=seen, stats=repo.user_stats(user["id"]))


@app.get("/content/{cid}", response_class=HTMLResponse)
def content_detail(request: Request, cid: int, channel: int = 0):
    user = auth.require_user(request)
    item = repo.get_content(cid)
    if not item:
        raise HTTPException(404, "Контент олдсонгүй.")
    view_id = repo.log_view(cid, user["id"], channel or None,
                            request.headers.get("user-agent", "")[:120])
    related = [r for r in repo.list_contents(section=item["section"], status="published", limit=7)
               if r["id"] != cid][:6]
    return page(request, "content.html", item=item, view_id=view_id,
                comments=repo.list_comments(cid), tags=repo.get_tags(cid),
                distributions=repo.list_distributions(cid, 20), related=related,
                my_views=repo.query_one(
                    "SELECT COUNT(*) AS n FROM views WHERE content_id=? AND user_id=?",
                    (cid, user["id"]))["n"])


@app.post("/content/{cid}/comment")
def post_comment(request: Request, cid: int, body: str = Form(...), rating: int = Form(None)):
    user = auth.require_user(request)
    if body.strip():
        repo.add_comment(cid, user["id"], body, rating)
        flash(request, "Сэтгэгдэл нэмэгдлээ.")
    return RedirectResponse(f"/content/{cid}#comments", status_code=303)


@app.post("/comment/{comment_id}/delete")
def remove_comment(request: Request, comment_id: int):
    user = auth.require_user(request)
    row = repo.query_one("SELECT * FROM comments WHERE id = ?", (comment_id,))
    if row and (row["user_id"] == user["id"] or user["role"] == ROLE_ADMIN):
        repo.delete_comment(comment_id)
        flash(request, "Сэтгэгдэл устгагдлаа.")
    return RedirectResponse(request.headers.get("referer", "/library"), status_code=303)


@app.post("/api/view/{view_id}/progress")
async def view_progress(request: Request, view_id: int):
    auth.require_user(request)
    data = await request.json()
    repo.update_view_progress(view_id, int(data.get("watched", 0)), bool(data.get("completed")))
    return {"ok": True}


# ------------------------------------------------------------- media
@app.get("/media/{cid}")
def media(request: Request, cid: int, download: int = 0):
    auth.require_user(request)
    item = repo.get_content(cid)
    if not item or not item["file_path"]:
        raise HTTPException(404, "Файл байршаагүй байна (demo контент).")
    return storage.serve(item["file_path"], request.headers.get("range"),
                         item["file_name"] if download else None)


@app.get("/thumb/{cid}")
def thumb(request: Request, cid: int):
    item = repo.get_content(cid)
    if not item or not item["thumb_path"]:
        raise HTTPException(404, "Cover зураг алга.")
    return storage.serve(item["thumb_path"])


# ============================================================ PROFILE
@app.get("/profile", response_class=HTMLResponse)
@app.get("/profile/{uid}", response_class=HTMLResponse)
def profile(request: Request, uid: int = 0):
    me = auth.require_user(request)
    target = me
    if uid and uid != me["id"]:
        if me["role"] != ROLE_ADMIN:
            raise HTTPException(403, "Өөр хэрэглэгчийн профайл харах эрхгүй.")
        target = repo.get_user(uid) or me
    stats = repo.user_stats(target["id"])
    return page(request, "profile.html", target=target, stats=stats,
                history=repo.user_history(target["id"]),
                my_comments=repo.user_comments(target["id"]),
                trend=json.dumps(repo.user_trend(target["id"], 30)),
                tier=tier_for(stats["views_total"]))


@app.post("/profile/update")
def profile_update(request: Request, position: str = Form(""), branch: str = Form(""),
                   department: str = Form(""), phone: str = Form("")):
    user = auth.require_user(request)
    repo.update_profile(user["id"], position, branch, department, phone)
    flash(request, "Профайл шинэчлэгдлээ.")
    return RedirectResponse("/profile", status_code=303)


# ========================================================== DASHBOARD
def _filters(request: Request) -> dict:
    qp = request.query_params
    return {
        "date_from": qp.get("date_from") or None,
        "date_to": qp.get("date_to") or None,
        "product_id": int(qp.get("product_id") or 0) or None,
        "channel_id": int(qp.get("channel_id") or 0) or None,
        "section": qp.get("section") or None,
    }


@app.get("/dashboard", response_class=HTMLResponse)
def dashboard(request: Request):
    auth.require_admin(request)
    f = _filters(request)
    days = int(request.query_params.get("days") or 30)
    kpi = repo.kpi_summary(**f)
    by_product = repo.views_by_product(10, **f)
    by_section = repo.views_by_section(**f)
    by_channel = repo.views_by_channel(**f)
    charts = {
        "trend": repo.views_trend(days, **f),
        "product": {"labels": [r["name"] for r in by_product],
                    "views": [r["views"] for r in by_product],
                    "viewers": [r["viewers"] for r in by_product],
                    "colors": [r["color"] for r in by_product]},
        "section": {"labels": [SECTION_SHORT.get(r["section"], r["section"]) for r in by_section],
                    "views": [r["views"] for r in by_section]},
        "channel": {"labels": [r["name"] for r in by_channel],
                    "views": [r["views"] for r in by_channel]},
        "hourly": repo.hourly_heat(**f),
        "matrix": repo.user_product_matrix(8, 6, **f),
    }
    return page(request, "dashboard.html", kpi=kpi, charts=json.dumps(charts),
                top_contents=repo.top_contents(8, **f), top_users=repo.top_users(10, **f),
                by_product=by_product, products=repo.list_products(),
                channels=repo.list_channels(), f=f, days=days,
                recent=repo.list_distributions(limit=8))


@app.get("/api/analytics")
def api_analytics(request: Request):
    auth.require_admin(request)
    f = _filters(request)
    days = int(request.query_params.get("days") or 30)
    return {
        "kpi": repo.kpi_summary(**f),
        "trend": repo.views_trend(days, **f),
        "products": [dict(r) for r in repo.views_by_product(10, **f)],
        "channels": [dict(r) for r in repo.views_by_channel(**f)],
        "top_contents": [dict(r) for r in repo.top_contents(8, **f)],
        "top_users": [dict(r) for r in repo.top_users(10, **f)],
    }


@app.get("/dashboard/export.csv")
def export_csv(request: Request):
    auth.require_admin(request)
    rows = repo.export_rows(**_filters(request))
    buf = io.StringIO()
    buf.write("﻿")  # Excel-д кирилл зөв харагдахын тулд BOM
    w = csv.writer(buf)
    w.writerow(["Огноо", "Хэрэглэгч", "Мэйл", "Албан тушаал", "Салбар", "Контент",
                "Хэсэг", "Төрөл", "Бүтээгдэхүүн", "Channel", "Үзсэн (сек)", "Дуустал", "Төхөөрөмж"])
    for r in rows:
        w.writerow([r["viewed_at"], r["user_name"], r["email"], r["position"], r["branch"],
                    r["content_title"], SECTION_SHORT.get(r["section"], r["section"]),
                    KIND_LABELS.get(r["kind"], r["kind"]), r["product"], r["channel"],
                    r["watched_sec"], "Тийм" if r["completed"] else "Үгүй", r["device"]])
    buf.seek(0)
    fname = f"mp-team-report-{datetime.now():%Y%m%d-%H%M}.csv"
    return StreamingResponse(iter([buf.getvalue()]), media_type="text/csv; charset=utf-8",
                             headers={"content-disposition": f'attachment; filename="{fname}"'})


# ====================================================== ADMIN: CONTENT
@app.get("/admin/content", response_class=HTMLResponse)
def admin_content(request: Request, section: str = "", q: str = "", status: str = ""):
    auth.require_admin(request)
    return page(request, "admin_content.html",
                items=repo.list_contents(section=section or None, search=q or None,
                                         status=status or None),
                products=repo.list_products(), channels=repo.list_channels(active_only=True),
                filters={"section": section, "q": q, "status": status})


@app.get("/admin/content/new", response_class=HTMLResponse)
def admin_content_new(request: Request):
    auth.require_admin(request)
    return page(request, "admin_upload.html", products=repo.list_products(),
                channels=repo.list_channels(active_only=True), item=None)


@app.post("/admin/content/new")
async def admin_content_create(
    request: Request,
    title: str = Form(...),
    description: str = Form(""),
    section: str = Form(SECTION_OTHER),
    product_id: str = Form(""),
    publish_date: str = Form(""),
    tags: str = Form(""),
    status: str = Form("published"),
    file: UploadFile = File(None),
    cover: UploadFile = File(None),
    channel_ids: list[str] = Form([]),
    message: str = Form(""),
    scheduled_at: str = Form(""),
):
    user = auth.require_admin(request)
    meta = {"file_name": None, "file_path": None, "mime": None, "size_bytes": 0, "kind": "poster"}
    if file is not None and file.filename:
        meta = storage.save_upload(file)
    thumb_path = None
    if cover is not None and cover.filename:
        thumb_path = storage.save_upload(cover, "thumbs")["file_path"]
    pid = int(product_id) if product_id else None
    color = "#00e0a4"
    if pid:
        p = repo.query_one("SELECT color FROM products WHERE id=?", (pid,))
        color = p["color"] if p else color
    if not thumb_path:
        if meta["kind"] == "image" and meta["file_path"]:
            thumb_path = meta["file_path"]
        else:
            thumb_path = storage.make_placeholder_thumb(title, KIND_LABELS.get(meta["kind"], ""), color)

    cid = repo.create_content(
        title=title.strip(), description=description.strip(), section=section,
        kind=meta["kind"], product_id=pid, file_name=meta["file_name"],
        file_path=meta["file_path"], mime=meta["mime"], size_bytes=meta["size_bytes"],
        thumb_path=thumb_path, duration_sec=0,
        publish_date=publish_date or datetime.now().date().isoformat(),
        status=status, created_by=user["id"],
    )
    repo.set_tags(cid, [t for t in tags.replace("#", "").split(",")])
    chans = [int(c) for c in channel_ids if str(c).strip().isdigit()]
    if chans:
        n = repo.distribute(cid, chans, message, scheduled_at or None, user["id"])
        flash(request, f"Контент хадгалагдаж, {n} channel руу "
                       f"{'төлөвлөгдлөө' if scheduled_at else 'илгээгдлээ'}.")
    else:
        flash(request, "Контент хадгалагдлаа.")
    return RedirectResponse(f"/admin/content", status_code=303)


@app.get("/admin/content/{cid}/edit", response_class=HTMLResponse)
def admin_content_edit(request: Request, cid: int):
    auth.require_admin(request)
    item = repo.get_content(cid)
    if not item:
        raise HTTPException(404, "Контент олдсонгүй.")
    return page(request, "admin_upload.html", products=repo.list_products(),
                channels=repo.list_channels(active_only=True), item=item,
                tags=", ".join(repo.get_tags(cid)),
                distributions=repo.list_distributions(cid))


@app.post("/admin/content/{cid}/edit")
async def admin_content_update(request: Request, cid: int, title: str = Form(...),
                               description: str = Form(""), section: str = Form(SECTION_OTHER),
                               product_id: str = Form(""), publish_date: str = Form(""),
                               tags: str = Form(""), status: str = Form("published"),
                               cover: UploadFile = File(None)):
    auth.require_admin(request)
    item = repo.get_content(cid)
    if not item:
        raise HTTPException(404, "Контент олдсонгүй.")
    fields = {
        "title": title.strip(), "description": description.strip(), "section": section,
        "product_id": int(product_id) if product_id else None,
        "publish_date": publish_date or item["publish_date"], "status": status,
    }
    if cover is not None and cover.filename:
        new_thumb = storage.save_upload(cover, "thumbs")["file_path"]
        if item["thumb_path"] and item["thumb_path"] != item["file_path"]:
            storage.delete_file(item["thumb_path"])
        fields["thumb_path"] = new_thumb
    repo.update_content(cid, **fields)
    repo.set_tags(cid, [t for t in tags.replace("#", "").split(",")])
    flash(request, "Контент шинэчлэгдлээ.")
    return RedirectResponse(f"/admin/content/{cid}/edit", status_code=303)


@app.post("/admin/content/{cid}/delete")
def admin_content_delete(request: Request, cid: int):
    auth.require_admin(request)
    item = repo.get_content(cid)
    if item:
        storage.delete_file(item["file_path"])
        if item["thumb_path"] != item["file_path"]:
            storage.delete_file(item["thumb_path"])
        repo.delete_content(cid)
        flash(request, "Контент устгагдлаа.", "warn")
    return RedirectResponse("/admin/content", status_code=303)


# ================================================= ADMIN: DISTRIBUTION
@app.get("/admin/distribute", response_class=HTMLResponse)
def admin_distribute_page(request: Request, content_id: int = 0):
    auth.require_admin(request)
    return page(request, "admin_distribute.html",
                contents=repo.list_contents(), channels=repo.list_channels(),
                selected=content_id, history=repo.list_distributions(limit=60))


@app.post("/admin/distribute")
def admin_distribute(request: Request, content_id: int = Form(...),
                     channel_ids: list[str] = Form([]), message: str = Form(""),
                     scheduled_at: str = Form("")):
    user = auth.require_admin(request)
    chans = [int(c) for c in channel_ids if str(c).strip().isdigit()]
    if not chans:
        flash(request, "Дор хаяж нэг Teams channel сонгоно уу.", "error")
        return RedirectResponse(f"/admin/distribute?content_id={content_id}", status_code=303)
    n = repo.distribute(content_id, chans, message, scheduled_at or None, user["id"])
    flash(request, f"{n} channel руу {'төлөвлөгдлөө' if scheduled_at else 'илгээгдлээ'}.")
    return RedirectResponse("/admin/distribute", status_code=303)


@app.post("/admin/distribute/{dist_id}/send-now")
def distribute_send_now(request: Request, dist_id: int):
    auth.require_admin(request)
    repo.mark_scheduled_sent(dist_id)
    flash(request, "Илгээгдлээ.")
    return RedirectResponse(request.headers.get("referer", "/admin/distribute"), status_code=303)


# ==================================================== ADMIN: CHANNELS
@app.get("/admin/channels", response_class=HTMLResponse)
def admin_channels(request: Request):
    auth.require_admin(request)
    return page(request, "admin_channels.html", channels=repo.list_channels())


@app.post("/admin/channels")
def admin_channels_add(request: Request, name: str = Form(...), team_name: str = Form(""),
                       link: str = Form(...), webhook_url: str = Form(""),
                       members: int = Form(0)):
    auth.require_admin(request)
    repo.add_channel(name, team_name, link, webhook_url, members)
    flash(request, "Teams channel нэмэгдлээ.")
    return RedirectResponse("/admin/channels", status_code=303)


@app.post("/admin/channels/{cid}/toggle")
def admin_channel_toggle(request: Request, cid: int):
    auth.require_admin(request)
    repo.toggle_channel(cid)
    return RedirectResponse("/admin/channels", status_code=303)


@app.post("/admin/channels/{cid}/delete")
def admin_channel_delete(request: Request, cid: int):
    auth.require_admin(request)
    repo.delete_channel(cid)
    flash(request, "Channel устгагдлаа.", "warn")
    return RedirectResponse("/admin/channels", status_code=303)


# ==================================================== ADMIN: PRODUCTS
@app.get("/admin/products", response_class=HTMLResponse)
def admin_products(request: Request):
    auth.require_admin(request)
    return page(request, "admin_products.html", products=repo.list_products())


@app.post("/admin/products")
def admin_products_add(request: Request, name: str = Form(...), code: str = Form(""),
                       line: str = Form("")):
    auth.require_admin(request)
    if repo.add_product(name, code, line):
        flash(request, f"'{name}' бүтээгдэхүүн нэмэгдлээ.")
    else:
        flash(request, "Ийм нэртэй бүтээгдэхүүн аль хэдийн бүртгэлтэй байна.", "warn")
    return RedirectResponse("/admin/products", status_code=303)


@app.post("/admin/products/import")
async def admin_products_import(request: Request, file: UploadFile = File(...)):
    """Excel (.xlsx) эсвэл CSV файлаас бүтээгдэхүүний нэрсийг нэг дор оруулах.

    Багана: 1-р багана = нэр (заавал), 2-р = код, 3-р = чиглэл. Толгой мөрийг таньж алгасна."""
    auth.require_admin(request)
    raw = await file.read()
    rows: list[list[str]] = []
    fname = (file.filename or "").lower()
    try:
        if fname.endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            for r in wb.active.iter_rows(values_only=True):
                rows.append(["" if c is None else str(c).strip() for c in r])
        else:
            text = raw.decode("utf-8-sig", errors="replace")
            delim = ";" if text.count(";") > text.count(",") else ","
            rows = [r for r in csv.reader(io.StringIO(text), delimiter=delim)]
    except Exception as e:
        flash(request, f"Файл уншиж чадсангүй: {e}", "error")
        return RedirectResponse("/admin/products", status_code=303)

    added, skipped = 0, 0
    for i, r in enumerate(rows):
        if not r or not (r[0] or "").strip():
            continue
        first = r[0].strip()
        if i == 0 and first.lower() in {"нэр", "name", "бүтээгдэхүүн", "product", "бүтээгдэхүүний нэр"}:
            continue
        if repo.add_product(first, r[1] if len(r) > 1 else None, r[2] if len(r) > 2 else None):
            added += 1
        else:
            skipped += 1
    flash(request, f"Импорт дууслаа: {added} шинэ, {skipped} давхардсан/алгассан.",
          "ok" if added else "warn")
    return RedirectResponse("/admin/products", status_code=303)


@app.get("/admin/products/template.csv")
def products_template(request: Request):
    auth.require_admin(request)
    body = "﻿Нэр,Код,Чиглэл\nАмпилин,AMP,Антибиотик\nГепаклин,GEP,Элэг\n"
    return PlainTextResponse(body, media_type="text/csv; charset=utf-8",
                             headers={"content-disposition": 'attachment; filename="products-template.csv"'})


@app.post("/admin/products/{pid}/edit")
def admin_product_edit(request: Request, pid: int, name: str = Form(...),
                       code: str = Form(""), line: str = Form(""), color: str = Form("")):
    auth.require_admin(request)
    if not repo.get_product(pid):
        raise HTTPException(404, "Бүтээгдэхүүн олдсонгүй.")
    if repo.update_product(pid, name, code, line, color):
        flash(request, f"'{name.strip()}' шинэчлэгдлээ.")
    else:
        flash(request, "Нэр хоосон эсвэл өөр бүтээгдэхүүнтэй давхардаж байна.", "error")
    return RedirectResponse("/admin/products", status_code=303)


@app.post("/admin/products/{pid}/delete")
def admin_product_delete(request: Request, pid: int):
    auth.require_admin(request)
    repo.delete_product(pid)
    flash(request, "Бүтээгдэхүүн устгагдлаа.", "warn")
    return RedirectResponse("/admin/products", status_code=303)


# ======================================================= ADMIN: USERS
@app.get("/admin/users", response_class=HTMLResponse)
def admin_users(request: Request):
    auth.require_admin(request)
    return page(request, "admin_users.html", users=repo.list_users())


@app.post("/admin/users/{uid}/role")
def admin_user_role(request: Request, uid: int, role: str = Form(...)):
    auth.require_admin(request)
    repo.update_user_role(uid, role if role in ROLE_LABELS else ROLE_VIEWER)
    flash(request, "Эрх шинэчлэгдлээ.")
    return RedirectResponse("/admin/users", status_code=303)


@app.get("/healthz")
def healthz():
    return {"status": "ok", "app": APP_NAME, "time": datetime.now().isoformat()}
